from __future__ import annotations

import hashlib
import json
import logging
import shutil
import traceback
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from time import perf_counter
from typing import Any, Iterator

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .config import AppConfig, load_business_context, load_column_mappings, load_human_tests, load_table_pairs
from .context_store import make_context_retriever
from .reporting import write_json


APPROVAL_COLUMNS = [
    "item_id", "item_type", "pair_id", "table_name", "column_name",
    "why_approval_required", "existing_value", "proposed_value", "evidence",
    "confidence", "user_action_required", "approval_status", "reviewer_comments",
    "reviewed_by", "reviewed_at",
]
APPROVAL_STATUSES = {"PENDING", "APPROVE", "REJECT", "NEEDS_CHANGES", "OVERRIDE"}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def canonical_json(value: Any) -> str:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), default=str)


def stable_id(*parts: Any, length: int = 32) -> str:
    raw = "|".join(str(part) for part in parts).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()[:length]


def payload_hash(payload: dict[str, Any]) -> str:
    return hashlib.sha256(canonical_json(payload).encode("utf-8")).hexdigest()


def workflow_paths(config: AppConfig, run_id: str, workflow: str) -> dict[str, Path]:
    approval_root = config.path(config.project.approvals.root)
    paths = {
        "approval_root": approval_root,
        "pending": approval_root / "pending",
        "reviewed": approval_root / "reviewed",
        "processed": approval_root / "processed",
        "archive": approval_root / "archive",
        "output": config.path(config.project.outputs_dir) / run_id,
        "logs": config.path(config.project.logs_dir) / run_id / workflow,
    }
    for path in paths.values():
        path.mkdir(parents=True, exist_ok=True)
    paths["log"] = paths["logs"] / "workflow.log"
    paths["checkpoint"] = paths["logs"] / "checkpoint.json"
    return paths


def context_workflow_paths(config: AppConfig, run_id: str) -> dict[str, Path]:
    return workflow_paths(config, run_id, "context")


def configure_workflow_logging(log_path: Path, level: str = "INFO") -> logging.Logger:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger(f"dq_agent.workflow.{log_path.parent.name}.{log_path.stem}")
    logger.setLevel(getattr(logging, level.upper(), logging.INFO))
    logger.propagate = False
    if logger.handlers:
        return logger
    formatter = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(name)s | %(funcName)s:%(lineno)d | %(message)s"
    )
    file_handler = logging.FileHandler(log_path, encoding="utf-8")
    file_handler.setFormatter(formatter)
    stream_handler = logging.StreamHandler()
    stream_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    logger.addHandler(stream_handler)
    return logger


@contextmanager
def logged_step(
    logger: logging.Logger, checkpoint_path: Path, step: str, **details: Any
) -> Iterator[None]:
    started = perf_counter()
    logger.info("%s STARTED | %s", step, details)
    try:
        yield
    except Exception as exc:
        checkpoint = _read_json(checkpoint_path)
        checkpoint.update({
            "failed_step": step, "failed_at": utc_now(), "error": str(exc),
            "stack_trace": traceback.format_exc(),
        })
        write_json(checkpoint_path, checkpoint)
        logger.exception("%s FAILED | %s", step, details)
        raise
    checkpoint = _read_json(checkpoint_path)
    checkpoint.update({
        "last_completed_step": step, "completed_at": utc_now(),
        "failed_step": None, "error": None, "stack_trace": None,
    })
    write_json(checkpoint_path, checkpoint)
    logger.info("%s COMPLETED | elapsed_seconds=%.3f | %s", step, perf_counter() - started, details)


def _read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def make_context_record(
    config: AppConfig,
    run_id: str,
    context_type: str,
    subject_key: str,
    payload: dict[str, Any],
    source_file: str,
    **identity: Any,
) -> dict[str, Any]:
    del config
    return {
        "context_type": context_type,
        "subject_key": subject_key,
        "payload": payload,
        "pair_id": identity.get("pair_id"),
        "target_table": identity.get("target_table"),
        "source_table": identity.get("source_table"),
        "column_name": identity.get("column_name"),
        "business_entity": identity.get("business_entity"),
        "table_type": identity.get("table_type"),
        "confidence": float(identity.get("confidence", 1.0)),
        "origin": identity.get("origin", "AGENT_INFERENCE"),
        "source_file": source_file,
        "run_id": run_id,
        "evidence": identity.get("evidence", []),
        "content_hash": payload_hash(payload),
    }


def normalize_input_context(
    config: AppConfig, run_id: str, logger: logging.Logger | None = None
) -> pd.DataFrame:
    records: list[dict[str, Any]] = []
    for pair in load_table_pairs(config):
        records.append(make_context_record(
            config, run_id, "runtime_table", pair.pair_id, pair.model_dump(),
            config.project.table_mappings, pair_id=pair.pair_id,
            target_table=pair.target_name, source_table=pair.source_name, origin="RUNTIME_INPUT",
        ))
    for mapping in load_column_mappings(config):
        records.append(make_context_record(
            config, run_id, "runtime_mapping", f"{mapping.pair_id}:{mapping.source_column}",
            mapping.model_dump(), config.project.column_mappings, pair_id=mapping.pair_id,
            column_name=mapping.source_column, origin="RUNTIME_INPUT",
        ))
    for test in load_human_tests(config):
        records.append(make_context_record(
            config, run_id, "human_test", test.id, test.model_dump(by_alias=False),
            config.project.human_tests, pair_id=test.pair_id, origin="RUNTIME_INPUT",
        ))
    for table_id, payload in (load_business_context(config).get("tables") or {}).items():
        records.append(make_context_record(
            config, run_id, "table", str(table_id), dict(payload or {}),
            config.project.context_tables, pair_id=str(table_id), origin="DEVELOPER_CONTEXT",
        ))
    if logger:
        logger.info("NORMALIZE_CONTEXT records=%s", len(records))
    return pd.DataFrame(records)


def validate_normalized_context(records: pd.DataFrame) -> None:
    if records.empty:
        return
    required = {"context_type", "subject_key", "payload", "content_hash"}
    missing = required - set(records.columns)
    if missing:
        raise ValueError(f"Normalized context is missing columns: {sorted(missing)}")


def build_context_proposals(
    records: pd.DataFrame, trusted: pd.DataFrame, run_id: str
) -> pd.DataFrame:
    del trusted
    proposals: list[dict[str, Any]] = []
    for row in records.to_dict("records"):
        if str(row.get("origin")) not in {"AGENT_INFERENCE", "LLM_INFERENCE"}:
            continue
        proposals.append({
            "item_id": stable_id(run_id, row.get("context_type"), row.get("subject_key"), row.get("content_hash")),
            "item_type": row.get("context_type"), "pair_id": row.get("pair_id"),
            "table_name": row.get("target_table"), "column_name": row.get("column_name"),
            "why_approval_required": "New reusable agent inference",
            "existing_value": "", "proposed_value": canonical_json(row.get("payload") or {}),
            "evidence": canonical_json(row.get("evidence") or {}),
            "confidence": row.get("confidence", 0),
            "user_action_required": "Approve, reject, or override the proposed value",
            "approval_status": "PENDING", "reviewer_comments": "", "reviewed_by": "",
            "reviewed_at": "", "context_record": row,
        })
    return pd.DataFrame(proposals)


def rank_relevant_context(
    current: pd.DataFrame, trusted: pd.DataFrame, threshold: float = 0.0, limit: int = 8
) -> pd.DataFrame:
    del current, threshold
    if trusted.empty:
        return trusted
    score = "retrieval_score" if "retrieval_score" in trusted.columns else "confidence"
    return trusted.sort_values(score, ascending=False).head(limit).reset_index(drop=True)


def _proposal_value(row: dict[str, Any]) -> str:
    value = row.get("proposed_value", row.get("proposal", row.get("payload", "")))
    return value if isinstance(value, str) else canonical_json(value)


def write_approval_workbook(path: Path, proposals: pd.DataFrame) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    rows: list[dict[str, Any]] = []
    for proposal in proposals.to_dict("records"):
        rows.append({
            "item_id": proposal.get("item_id") or proposal.get("inference_id") or stable_id(_proposal_value(proposal)),
            "item_type": proposal.get("item_type") or proposal.get("category"),
            "pair_id": proposal.get("pair_id"), "table_name": proposal.get("table_name"),
            "column_name": proposal.get("column_name") or proposal.get("subject"),
            "why_approval_required": proposal.get("why_approval_required") or "Agent inference requires review",
            "existing_value": proposal.get("existing_value", ""),
            "proposed_value": _proposal_value(proposal),
            "evidence": proposal.get("evidence", ""), "confidence": proposal.get("confidence", 0),
            "user_action_required": proposal.get("user_action_required") or "Approve, reject, or override",
            "approval_status": "PENDING", "reviewer_comments": "", "reviewed_by": "", "reviewed_at": "",
        })
    pd.DataFrame(rows, columns=APPROVAL_COLUMNS).to_excel(path, index=False, sheet_name="Approvals")
    workbook = load_workbook(path)
    sheet = workbook["Approvals"]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [24, 20, 20, 34, 22, 42, 40, 50, 45, 12, 38, 18, 40, 22, 22]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    validation = DataValidation(type="list", formula1='"PENDING,APPROVE,REJECT,NEEDS_CHANGES,OVERRIDE"', allow_blank=False)
    sheet.add_data_validation(validation)
    validation.add(f"L2:L{max(sheet.max_row, 2)}")
    sheet.conditional_formatting.add(
        f"L2:L{max(sheet.max_row, 2)}",
        FormulaRule(formula=['$L2="REJECT"'], fill=PatternFill("solid", fgColor="F4CCCC")),
    )
    workbook.save(path)
    return path


def read_approval_workbook(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Approval workbook not found: {path}")
    frame = pd.read_excel(path, dtype=object).where(lambda value: pd.notna(value), "")
    missing = set(APPROVAL_COLUMNS) - set(frame.columns)
    if missing:
        raise ValueError(f"Approval workbook is missing columns: {sorted(missing)}")
    frame["approval_status"] = frame["approval_status"].astype(str).str.strip().str.upper()
    invalid = sorted(set(frame["approval_status"]) - APPROVAL_STATUSES)
    if invalid:
        raise ValueError(f"Approval workbook contains invalid statuses: {invalid}")
    return frame


def approval_rows_to_context(frame: pd.DataFrame) -> list[dict[str, Any]]:
    approved: list[dict[str, Any]] = []
    for row in frame.to_dict("records"):
        if row["approval_status"] not in {"APPROVE", "OVERRIDE"}:
            continue
        proposed = row["proposed_value"]
        try:
            payload = json.loads(str(proposed))
        except json.JSONDecodeError:
            payload = {"value": proposed}
        approved.append({
            "context_type": row["item_type"], "subject_key": row["item_id"],
            "pair_id": row.get("pair_id") or None, "target_table": row.get("table_name") or None,
            "column_name": row.get("column_name") or None, "payload": payload,
            "confidence": float(row.get("confidence") or 0), "reviewed_by": row.get("reviewed_by") or None,
            "reviewer_comments": row.get("reviewer_comments") or None,
        })
    return approved


def process_approval_workbook(config: AppConfig, path: Path) -> dict[str, Any]:
    reviewed_dir = (config.path(config.project.approvals.root) / "reviewed").resolve()
    if path.resolve().parent != reviewed_dir:
        raise ValueError(f"Approval workbook must be placed under {reviewed_dir} before processing")
    frame = read_approval_workbook(path)
    if (frame["approval_status"] == "PENDING").any():
        raise ValueError("Every approval row must be reviewed before processing")
    approved = approval_rows_to_context(frame)
    written = make_context_retriever(config).upsert_approved(approved)
    destination = archive_approval_workbook(path, config.path(config.project.approvals.root) / "processed")
    return {"approved": written, "processed_file": str(destination)}


def archive_approval_workbook(source: Path, destination_dir: Path) -> Path:
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / source.name
    if destination.exists():
        archive = destination_dir.parent / "archive"
        archive.mkdir(parents=True, exist_ok=True)
        shutil.move(str(destination), str(archive / f"{destination.stem}_{stable_id(utc_now(), length=8)}{destination.suffix}"))
    return Path(shutil.move(str(source), str(destination)))
