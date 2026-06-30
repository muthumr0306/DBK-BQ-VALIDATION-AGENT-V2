from __future__ import annotations

import copy
import json
import re
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from .config import (
    ColumnMapping,
    HumanTest,
    LLMConfig,
    TablePair,
    load_app_config,
    read_yaml,
)
from .context_store import make_context_retriever
from .context_utils import (
    APPROVAL_COLUMNS,
    APPROVAL_STATUSES,
    approval_rows_to_context,
    process_approval_workbook,
    read_approval_workbook,
    stable_id,
    utc_now,
)
from .workflow import DQWorkflow


TABLE_MAPPING_COLUMNS = [
    "pair_id", "enabled", "mode", "source_catalog", "source_schema", "source_table",
    "target_project", "target_dataset", "target_table",
]
COLUMN_MAPPING_COLUMNS = ["pair_id", "source_column", "target_column", "status"]
FILTER_COLUMNS = ["enabled", "pair_id", "kind", "side", "column", "operator", "value"]
CONTEXT_HINT_COLUMNS = [
    "pair_id", "table_type", "business_entity", "primary_key_source", "primary_key_target",
    "audit_column_source", "audit_column_target", "freshness_sla_minutes", "scd2_enabled",
    "description",
]
HUMAN_TEST_COLUMNS = [
    "id", "pair_id", "enabled", "type", "scope", "column", "columns", "source_column",
    "target_column", "source_columns", "target_columns", "values", "allowed_nulls",
    "allowed_duplicates", "max_delay_minutes", "aggregation", "comparison", "group_by",
    "expected_value", "tolerance_absolute", "tolerance_percentage", "severity", "description",
    "source_sql", "target_sql",
]
RESULT_SHEETS = [
    "Schema", "Row Counts", "Freshness", "Profiles", "Business Rules", "Human Tests",
    "Relationships", "Measures", "Failures",
]
UI_CATEGORIES = [
    "Schema checks", "Row count checks", "Null checks", "Duplicate checks",
    "Freshness checks", "Referential integrity checks", "Business rule checks",
    "Measure reconciliation checks",
]


@dataclass(frozen=True)
class UIRunFiles:
    run_id: str
    run_dir: Path
    ui_inputs_dir: Path
    project_file: str
    llm_file: str
    table_mappings: Path
    column_mappings: Path
    runtime_overrides: Path
    human_tests: Path
    context_tables: Path


def root_path(root: str | Path = ".") -> Path:
    return Path(root).resolve()


def utc_run_id(prefix: str = "ui_validation") -> str:
    return f"{prefix}_{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}"


def _clean(value: Any) -> Any:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def _clean_record(record: dict[str, Any]) -> dict[str, Any]:
    return {key: _clean(value) for key, value in record.items()}


def _as_bool(value: Any, default: bool = False) -> bool:
    value = _clean(value)
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "y", "enabled", "on"}


def _parse_scalar(value: Any) -> Any:
    value = _clean(value)
    if value is None:
        return None
    if isinstance(value, (bool, int, float)):
        return value
    text = str(value).strip()
    lowered = text.lower()
    if lowered in {"true", "false"}:
        return lowered == "true"
    if lowered in {"null", "none"}:
        return None
    try:
        if "." in text:
            return float(text)
        return int(text)
    except ValueError:
        return text


def _parse_list(value: Any) -> list[Any]:
    value = _clean(value)
    if value is None:
        return []
    if isinstance(value, list):
        return value
    return [_parse_scalar(item) for item in str(value).split(",") if str(item).strip()]


def _list_text(value: Any) -> str:
    if isinstance(value, list):
        return ", ".join(str(item) for item in value)
    return "" if value is None else str(value)


def _relative(root: Path, path: Path) -> str:
    return str(path.resolve().relative_to(root.resolve())).replace("\\", "/")


def _json_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, indent=2, default=str)


def _write_yaml(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(payload, sort_keys=False, allow_unicode=False), encoding="utf-8")


def _empty_frame(columns: list[str]) -> pd.DataFrame:
    return pd.DataFrame(columns=columns)


def _ensure_columns(frame: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    output = frame.copy() if frame is not None else pd.DataFrame()
    for column in columns:
        if column not in output.columns:
            output[column] = None
    return output[columns]


def read_tabular_upload(uploaded: Any) -> pd.DataFrame:
    name = str(getattr(uploaded, "name", "")).lower()
    if name.endswith((".xlsx", ".xls")):
        return pd.read_excel(uploaded, dtype=object)
    if name.endswith(".csv"):
        return pd.read_csv(uploaded, dtype=object)
    raise ValueError("Upload an Excel or CSV file.")


def read_yaml_upload(uploaded: Any) -> dict[str, Any]:
    return yaml.safe_load(uploaded.getvalue().decode("utf-8")) or {}


def load_runtime_inputs(root: str | Path = ".") -> dict[str, Any]:
    project_root = root_path(root)
    config = load_app_config(project_root)
    tables_path = config.path(config.project.table_mappings)
    columns_path = config.path(config.project.column_mappings)
    tables = pd.read_excel(tables_path, dtype=object) if tables_path.exists() else _empty_frame(TABLE_MAPPING_COLUMNS)
    columns = pd.read_excel(columns_path, dtype=object) if columns_path.exists() else _empty_frame(COLUMN_MAPPING_COLUMNS)
    runtime_overrides = read_yaml(config.path(config.project.runtime_overrides), default={"tables": {}})
    human_tests = read_yaml(config.path(config.project.human_tests), default={"tests": []})
    context_tables = read_yaml(config.path(config.project.context_tables), default={"tables": {}})
    return {
        "config": config,
        "table_mappings": _ensure_columns(tables, TABLE_MAPPING_COLUMNS),
        "column_mappings": _ensure_columns(columns, COLUMN_MAPPING_COLUMNS),
        "filter_rows": filters_to_frame(runtime_overrides, context_tables),
        "human_tests": human_tests_to_frame(human_tests.get("tests", [])),
        "context_hints": context_hints_to_frame(context_tables),
        "llm": config.llm.model_dump(),
        "context_tables": context_tables,
    }


def filters_to_frame(runtime_overrides: dict[str, Any], context_tables: dict[str, Any] | None = None) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for pair_id, payload in (runtime_overrides.get("tables") or {}).items():
        filters = (payload or {}).get("filters") or {}
        for side in ("source", "target"):
            for item in filters.get(side) or []:
                rows.append({
                    "enabled": True, "pair_id": pair_id, "kind": "Runtime filter", "side": side,
                    "column": item.get("column"), "operator": item.get("operator", "eq"),
                    "value": _list_text(item.get("value")) if item.get("operator") in {"in", "not_in"} else item.get("value"),
                })
    for pair_id, payload in ((context_tables or {}).get("tables") or {}).items():
        current = (((payload or {}).get("scd2") or {}).get("current_filter") or {})
        for side in ("source", "target"):
            for item in current.get(side) or []:
                rows.append({
                    "enabled": True, "pair_id": pair_id, "kind": "SCD current filter", "side": side,
                    "column": item.get("column"), "operator": item.get("operator", "eq"),
                    "value": _list_text(item.get("value")) if item.get("operator") in {"in", "not_in"} else item.get("value"),
                })
    return _ensure_columns(pd.DataFrame(rows), FILTER_COLUMNS)


def _filter_item(row: dict[str, Any]) -> dict[str, Any]:
    operator = str(row.get("operator") or "eq").strip().lower()
    item = {"column": row.get("column"), "operator": operator}
    if operator not in {"is_null", "not_null"}:
        item["value"] = _parse_list(row.get("value")) if operator in {"in", "not_in"} else _parse_scalar(row.get("value"))
    return item


def frame_to_runtime_overrides(frame: pd.DataFrame) -> dict[str, Any]:
    output: dict[str, Any] = {"tables": {}}
    for raw in _ensure_columns(frame, FILTER_COLUMNS).to_dict("records"):
        row = _clean_record(raw)
        if not _as_bool(row.get("enabled"), True) or not row.get("pair_id") or not row.get("column"):
            continue
        if row.get("kind") != "Runtime filter":
            continue
        side = str(row.get("side") or "target").lower()
        if side not in {"source", "target"}:
            continue
        table = output["tables"].setdefault(row["pair_id"], {"filters": {}})
        table["filters"].setdefault(side, []).append(_filter_item(row))
    return output


def human_tests_to_frame(records: list[dict[str, Any]]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for raw in records:
        row = dict(raw or {})
        row["id"] = row.get("id") or row.get("test_id")
        for column in ("columns", "source_columns", "target_columns", "values", "group_by"):
            row[column] = _list_text(row.get(column))
        rows.append(row)
    return _ensure_columns(pd.DataFrame(rows), HUMAN_TEST_COLUMNS)


def frame_to_human_tests(frame: pd.DataFrame) -> dict[str, Any]:
    tests: list[dict[str, Any]] = []
    for raw in _ensure_columns(frame, HUMAN_TEST_COLUMNS).to_dict("records"):
        row = _clean_record(raw)
        if not row.get("id") or not row.get("pair_id"):
            continue
        row["enabled"] = _as_bool(row.get("enabled"), True)
        for column in ("columns", "source_columns", "target_columns", "values", "group_by"):
            row[column] = _parse_list(row.get(column))
        for column in ("allowed_nulls", "allowed_duplicates", "max_delay_minutes"):
            if row.get(column) is not None:
                row[column] = int(float(row[column]))
        for column in ("tolerance_absolute", "tolerance_percentage"):
            if row.get(column) is not None:
                row[column] = float(row[column])
        if row.get("expected_value") is not None:
            row["expected_value"] = _parse_scalar(row.get("expected_value"))
        tests.append({key: value for key, value in row.items() if value not in (None, "")})
    return {"tests": tests}


def context_hints_to_frame(context_tables: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for pair_id, payload in (context_tables.get("tables") or {}).items():
        payload = payload or {}
        primary_key = payload.get("primary_key") or {}
        audit_columns = payload.get("audit_columns") or {}
        rows.append({
            "pair_id": pair_id,
            "table_type": payload.get("table_type"),
            "business_entity": payload.get("business_entity"),
            "primary_key_source": _list_text(primary_key.get("source")),
            "primary_key_target": _list_text(primary_key.get("target")),
            "audit_column_source": audit_columns.get("source"),
            "audit_column_target": audit_columns.get("target"),
            "freshness_sla_minutes": payload.get("freshness_sla_minutes"),
            "scd2_enabled": bool((payload.get("scd2") or {}).get("enabled", False)),
            "description": payload.get("description"),
        })
    return _ensure_columns(pd.DataFrame(rows), CONTEXT_HINT_COLUMNS)


def frame_to_context_tables(base_context: dict[str, Any], hints: pd.DataFrame, filter_rows: pd.DataFrame) -> dict[str, Any]:
    document = copy.deepcopy(base_context or {"tables": {}})
    document.setdefault("tables", {})
    for raw in _ensure_columns(hints, CONTEXT_HINT_COLUMNS).to_dict("records"):
        row = _clean_record(raw)
        pair_id = row.get("pair_id")
        if not pair_id:
            continue
        payload = document["tables"].setdefault(pair_id, {})
        for column in ("table_type", "business_entity", "description"):
            if row.get(column) is not None:
                payload[column] = row[column]
        primary_key = payload.setdefault("primary_key", {})
        source_keys = _parse_list(row.get("primary_key_source"))
        target_keys = _parse_list(row.get("primary_key_target"))
        if source_keys:
            primary_key["source"] = [str(item) for item in source_keys]
        if target_keys:
            primary_key["target"] = [str(item) for item in target_keys]
        audit_columns = payload.setdefault("audit_columns", {})
        if row.get("audit_column_source"):
            audit_columns["source"] = row["audit_column_source"]
        if row.get("audit_column_target"):
            audit_columns["target"] = row["audit_column_target"]
        if row.get("freshness_sla_minutes") is not None:
            payload["freshness_sla_minutes"] = int(float(row["freshness_sla_minutes"]))
        scd2 = payload.setdefault("scd2", {})
        scd2["enabled"] = _as_bool(row.get("scd2_enabled"), False)
    scd_replacements: dict[str, set[str]] = {}
    for raw in _ensure_columns(filter_rows, FILTER_COLUMNS).to_dict("records"):
        row = _clean_record(raw)
        if _as_bool(row.get("enabled"), True) and row.get("kind") == "SCD current filter" and row.get("pair_id"):
            side = str(row.get("side") or "target").lower()
            if side in {"source", "target"}:
                scd_replacements.setdefault(str(row["pair_id"]), set()).add(side)
    for pair_id, sides in scd_replacements.items():
        current = document["tables"].setdefault(pair_id, {}).setdefault("scd2", {}).setdefault("current_filter", {})
        for side in sides:
            current[side] = []
    for raw in _ensure_columns(filter_rows, FILTER_COLUMNS).to_dict("records"):
        row = _clean_record(raw)
        if not _as_bool(row.get("enabled"), True) or row.get("kind") != "SCD current filter":
            continue
        if not row.get("pair_id") or not row.get("column"):
            continue
        side = str(row.get("side") or "target").lower()
        if side not in {"source", "target"}:
            continue
        table = document["tables"].setdefault(row["pair_id"], {})
        scd2 = table.setdefault("scd2", {})
        scd2["enabled"] = True
        scd2.setdefault("current_filter", {}).setdefault(side, []).append(_filter_item(row))
    return document


def _validate_table_rows(frame: pd.DataFrame) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    errors: list[str] = []
    for index, raw in enumerate(_ensure_columns(frame, TABLE_MAPPING_COLUMNS).to_dict("records"), start=2):
        row = _clean_record(raw)
        row["enabled"] = _as_bool(row.get("enabled"), True)
        if not row.get("pair_id"):
            continue
        try:
            TablePair.model_validate(row)
        except ValidationError as exc:
            errors.append(f"table_mappings row {index}: {exc.errors()[0]['msg']}")
        records.append(row)
    if errors:
        raise ValueError("; ".join(errors))
    return records


def _validate_column_rows(frame: pd.DataFrame) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    errors: list[str] = []
    for index, raw in enumerate(_ensure_columns(frame, COLUMN_MAPPING_COLUMNS).to_dict("records"), start=2):
        row = _clean_record(raw)
        if not row.get("pair_id") or not row.get("source_column") or not row.get("target_column"):
            continue
        try:
            ColumnMapping.model_validate(row)
        except ValidationError as exc:
            errors.append(f"column_mappings row {index}: {exc.errors()[0]['msg']}")
        records.append(row)
    if errors:
        raise ValueError("; ".join(errors))
    return records


def _validate_human_tests(document: dict[str, Any]) -> None:
    errors: list[str] = []
    for index, row in enumerate(document.get("tests") or [], start=2):
        try:
            HumanTest.model_validate(row)
        except ValidationError as exc:
            errors.append(f"human_tests row {index}: {exc.errors()[0]['msg']}")
    if errors:
        raise ValueError("; ".join(errors))


def prepare_ui_run(
    root: str | Path,
    run_id: str,
    table_mappings: pd.DataFrame,
    column_mappings: pd.DataFrame,
    filter_rows: pd.DataFrame,
    human_tests: pd.DataFrame,
    context_hints: pd.DataFrame,
    llm_overrides: dict[str, Any] | None = None,
) -> UIRunFiles:
    project_root = root_path(root)
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,127}", run_id):
        raise ValueError("Run ID may contain only letters, digits, dot, underscore, and hyphen.")
    config = load_app_config(project_root)
    run_dir = config.path(config.project.logs_dir) / run_id
    ui_dir = run_dir / "ui_inputs"
    ui_dir.mkdir(parents=True, exist_ok=True)

    table_records = _validate_table_rows(table_mappings)
    column_records = _validate_column_rows(column_mappings)
    human_test_document = frame_to_human_tests(human_tests)
    _validate_human_tests(human_test_document)
    runtime_overrides = frame_to_runtime_overrides(filter_rows)
    base_context = read_yaml(config.path(config.project.context_tables), default={"tables": {}})
    context_tables = frame_to_context_tables(base_context, context_hints, filter_rows)

    table_path = ui_dir / "table_mappings.xlsx"
    column_path = ui_dir / "column_mappings.xlsx"
    overrides_path = ui_dir / "runtime_overrides.yaml"
    tests_path = ui_dir / "human_tests.yaml"
    context_path = ui_dir / "tables.yaml"
    pd.DataFrame(table_records, columns=TABLE_MAPPING_COLUMNS).to_excel(table_path, index=False)
    pd.DataFrame(column_records, columns=COLUMN_MAPPING_COLUMNS).to_excel(column_path, index=False)
    _write_yaml(overrides_path, runtime_overrides)
    _write_yaml(tests_path, human_test_document)
    _write_yaml(context_path, context_tables)

    project_data = config.project.model_dump(mode="json")
    project_data.update({
        "table_mappings": _relative(project_root, table_path),
        "column_mappings": _relative(project_root, column_path),
        "runtime_overrides": _relative(project_root, overrides_path),
        "human_tests": _relative(project_root, tests_path),
        "context_tables": _relative(project_root, context_path),
    })
    project_file_path = run_dir / "ui_project.yaml"
    _write_yaml(project_file_path, project_data)

    llm_data = config.llm.model_dump(mode="json")
    for key, value in (llm_overrides or {}).items():
        if value not in (None, ""):
            llm_data[key] = value
    LLMConfig.model_validate(llm_data)
    llm_file_path = run_dir / "ui_llm.yaml"
    _write_yaml(llm_file_path, llm_data)

    return UIRunFiles(
        run_id=run_id,
        run_dir=run_dir,
        ui_inputs_dir=ui_dir,
        project_file=_relative(project_root, project_file_path),
        llm_file=_relative(project_root, llm_file_path),
        table_mappings=table_path,
        column_mappings=column_path,
        runtime_overrides=overrides_path,
        human_tests=tests_path,
        context_tables=context_path,
    )


def run_workflow_from_ui(
    root: str | Path,
    run_id: str,
    table_mappings: pd.DataFrame,
    column_mappings: pd.DataFrame,
    filter_rows: pd.DataFrame,
    human_tests: pd.DataFrame,
    context_hints: pd.DataFrame,
    llm_overrides: dict[str, Any] | None = None,
    stop_after: str = "reports",
) -> dict[str, Any]:
    files = prepare_ui_run(
        root, run_id, table_mappings, column_mappings, filter_rows,
        human_tests, context_hints, llm_overrides,
    )
    manifest = DQWorkflow(root_path(root), project_file=files.project_file, llm_file=files.llm_file).run(
        run_id=run_id, stop_after=stop_after
    )
    manifest["ui_project_file"] = files.project_file
    manifest["ui_inputs_dir"] = str(files.ui_inputs_dir)
    return manifest


def _read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def read_jsonl(path: Path, limit: int | None = None) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    rows: list[dict[str, Any]] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            rows.append(json.loads(line))
        except json.JSONDecodeError:
            rows.append({"message": line})
    return rows[-limit:] if limit else rows


def load_run_history(root: str | Path = ".") -> pd.DataFrame:
    project_root = root_path(root)
    config = load_app_config(project_root)
    rows: list[dict[str, Any]] = []
    logs_dir = config.path(config.project.logs_dir)
    for manifest_path in sorted(logs_dir.glob("*/manifest.json")):
        manifest = _read_json(manifest_path)
        if not manifest:
            continue
        run_id = manifest.get("run_id") or manifest_path.parent.name
        rows.append({
            "run_id": run_id,
            "status": manifest.get("status"),
            "started_at": manifest.get("started_at"),
            "completed_at": manifest.get("completed_at"),
            "tables": len(manifest.get("tables") or []),
            "approval_proposals": manifest.get("approval_proposals", 0),
            "report": str(config.path(config.project.outputs_dir) / run_id / "dq_validation_report.xlsx"),
        })
    frame = pd.DataFrame(rows)
    if frame.empty:
        return frame
    return frame.sort_values("started_at", ascending=False, na_position="last").reset_index(drop=True)


def latest_run_id(root: str | Path = ".") -> str | None:
    history = load_run_history(root)
    if history.empty:
        return None
    return str(history.iloc[0]["run_id"])


def _sheet_frame(frame: pd.DataFrame | None) -> pd.DataFrame:
    if frame is None or frame.empty:
        return pd.DataFrame()
    output = frame.where(pd.notna(frame), "")
    if list(output.columns) == ["message"] and len(output) == 1:
        return pd.DataFrame()
    return output


def load_report_workbook(root: str | Path, run_id: str) -> dict[str, pd.DataFrame]:
    config = load_app_config(root_path(root))
    workbook = config.path(config.project.outputs_dir) / run_id / "dq_validation_report.xlsx"
    if not workbook.exists():
        return {}
    sheets = pd.read_excel(workbook, sheet_name=None, dtype=object)
    return {name: _sheet_frame(frame) for name, frame in sheets.items()}


def consolidated_results(report: dict[str, pd.DataFrame]) -> pd.DataFrame:
    frames = []
    seen = set()
    for sheet in RESULT_SHEETS:
        frame = report.get(sheet)
        if frame is None or frame.empty:
            continue
        frame = frame.copy()
        frame["report_section"] = sheet
        for row in frame.to_dict("records"):
            key = row.get("rule_id") or json.dumps(row, sort_keys=True, default=str)
            if key in seen:
                continue
            seen.add(key)
            frames.append(row)
    return pd.DataFrame(frames)


def ui_category(row: dict[str, Any]) -> str:
    category = str(row.get("category") or "").lower()
    rule_type = str(row.get("type") or "").lower()
    if rule_type == "schema" or category == "schema":
        return "Schema checks"
    if category == "reconciliation" or rule_type in {"row_count", "key_values", "key_buckets", "row_reconciliation"}:
        return "Row count checks"
    if category == "completeness" or rule_type == "null_count":
        return "Null checks"
    if category == "uniqueness" or rule_type in {"uniqueness", "top_duplicates"}:
        return "Duplicate checks"
    if category == "freshness" or rule_type == "freshness":
        return "Freshness checks"
    if category == "referential_integrity" or rule_type == "relationship":
        return "Referential integrity checks"
    if rule_type == "measure_reconciliation" or category == "measure_reconciliation":
        return "Measure reconciliation checks"
    return "Business rule checks"


def _status(value: Any) -> str:
    text = str(value or "").upper()
    if text in {"PASS", "PASSED", "COMPLETED"}:
        return "Passed"
    if text in {"FAIL", "FAILED", "ERROR", "COMPLETED_WITH_FAILURES", "COMPLETED_WITH_ERRORS"}:
        return "Failed"
    if text in {"WAITING_FOR_REVIEW", "COMPLETED_WITH_REVIEWS", "COMPLETED_WAITING_FOR_REVIEW"}:
        return "Approval Required"
    if text in {"SKIP", "SKIPPED", "PENDING"}:
        return "Warning"
    return text.title() if text else "Unknown"


def _int_value(value: Any, default: int = 0) -> int:
    value = _clean(value)
    if value is None:
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def table_summary_for_ui(
    tables: pd.DataFrame,
    results: pd.DataFrame,
    rca: pd.DataFrame,
    approvals: pd.DataFrame,
) -> pd.DataFrame:
    columns = [
        "Table name", "Overall status", "Total checks", "Passed checks", "Failed checks",
        "Warning checks", "Approval required", "Main issue category", "RCA status",
        "Source table", "Target table", "Pair ID",
    ]
    if tables.empty:
        return pd.DataFrame(columns=columns)
    rows: list[dict[str, Any]] = []
    for raw in tables.to_dict("records"):
        pair_id = raw.get("pair_id") or raw.get("Pair ID")
        target = raw.get("target_table") or raw.get("Target table") or pair_id
        source = raw.get("source_table") or raw.get("Source table") or ""
        pair_results = results[results["pair_id"].astype(str) == str(pair_id)] if not results.empty and "pair_id" in results else pd.DataFrame()
        pair_approvals = approvals[approvals["pair_id"].astype(str) == str(pair_id)] if not approvals.empty and "pair_id" in approvals else pd.DataFrame()
        pair_rca = rca[rca["pair_id"].astype(str) == str(pair_id)] if not rca.empty and "pair_id" in rca else pd.DataFrame()
        statuses = pair_results.get("status", pd.Series(dtype=object)).astype(str).str.upper() if not pair_results.empty else pd.Series(dtype=object)
        failed_results = pair_results[statuses.isin(["FAIL", "ERROR"])] if not pair_results.empty else pd.DataFrame()
        issue_category = "None"
        if not failed_results.empty:
            counts = failed_results.apply(lambda row: ui_category(row.to_dict()), axis=1).value_counts()
            issue_category = str(counts.index[0])
        failed_count = _int_value(raw.get("failed"), int(statuses.isin(["FAIL", "ERROR"]).sum()))
        total_checks = _int_value(raw.get("rule_count"), len(pair_results))
        passed_checks = _int_value(raw.get("passed"), int((statuses == "PASS").sum()))
        warning_checks = int(statuses.isin(["SKIP", "PENDING", "WARNING", "WARN"]).sum())
        approval_count = _int_value(raw.get("review_items"), len(pair_approvals))
        rca_status = "Not required"
        if failed_count:
            rca_status = "RCA Complete" if not pair_rca.empty else "RCA Pending"
        rows.append({
            "Table name": str(target or pair_id),
            "Overall status": _status(raw.get("status")),
            "Total checks": total_checks,
            "Passed checks": passed_checks,
            "Failed checks": failed_count,
            "Warning checks": warning_checks,
            "Approval required": approval_count,
            "Main issue category": issue_category,
            "RCA status": rca_status,
            "Source table": source,
            "Target table": target,
            "Pair ID": pair_id,
        })
    return pd.DataFrame(rows, columns=columns)


def _demo_snapshot() -> dict[str, Any]:
    tables = pd.DataFrame([
        {"pair_id": "sales_fact", "source_table": "main.sales.fact_sales", "target_table": "prod.analytics.fact_sales", "status": "FAIL", "rule_count": 42, "passed": 36, "failed": 5, "review_items": 2},
        {"pair_id": "brand_dim", "source_table": "main.dimensions.dim_brand", "target_table": "prod.analytics.dim_brand", "status": "WAITING_FOR_REVIEW", "rule_count": 28, "passed": 26, "failed": 0, "review_items": 3},
        {"pair_id": "market_dim", "source_table": "main.dimensions.dim_market", "target_table": "prod.analytics.dim_market", "status": "PASS", "rule_count": 24, "passed": 24, "failed": 0, "review_items": 0},
        {"pair_id": "customer_dim", "source_table": "main.dimensions.dim_customer", "target_table": "prod.analytics.dim_customer", "status": "FAIL", "rule_count": 31, "passed": 28, "failed": 2, "review_items": 1},
    ])
    results = pd.DataFrame([
        {"pair_id": "sales_fact", "rule_id": "sales_fact__row_count", "category": "reconciliation", "type": "row_count", "status": "FAIL", "description": "Filtered source and target row counts match", "evidence": {"source": {"row_count": 1250340}, "target": {"row_count": 1249822}}, "comparison": {"difference": 518}},
        {"pair_id": "sales_fact", "rule_id": "sales_fact__freshness", "category": "freshness", "type": "freshness", "status": "FAIL", "description": "Target freshness is within SLA", "evidence": {"target": {"max_timestamp": "2026-06-28T06:00:00Z"}}, "comparison": {"allowed_minutes": 1440}},
        {"pair_id": "sales_fact", "rule_id": "sales_fact__relationship__brand", "category": "referential_integrity", "type": "relationship", "status": "FAIL", "description": "Sales records reference approved brand dimension rows", "evidence": {"target": {"orphan_count": 941}}, "comparison": {"checked_count": 1249822}},
        {"pair_id": "customer_dim", "rule_id": "customer_dim__target_key_unique", "category": "uniqueness", "type": "uniqueness", "status": "FAIL", "description": "Customer key is unique", "evidence": {"target": {"duplicate_groups": 19}}, "comparison": {"duplicate_rows": 21}},
        {"pair_id": "market_dim", "rule_id": "market_dim__freshness", "category": "freshness", "type": "freshness", "status": "PASS", "description": "Freshness is within SLA"},
    ])
    rca = pd.DataFrame([
        {"pair_id": "sales_fact", "rule_id": "sales_fact__row_count", "conclusion.classification": "likely", "conclusion.conclusion": "Target load excluded late-arriving ERP records because the runtime source_system filter was applied before the incremental cutoff.", "conclusion.confidence": 0.82, "diagnostics": "date coverage, filter impact"},
        {"pair_id": "sales_fact", "rule_id": "sales_fact__relationship__brand", "conclusion.classification": "confirmed", "conclusion.conclusion": "941 sales rows reference inactive brand dimension versions after SCD current-row filtering.", "conclusion.confidence": 0.91, "diagnostics": "relationship gap, inactive members"},
        {"pair_id": "customer_dim", "rule_id": "customer_dim__target_key_unique", "conclusion.classification": "possible", "conclusion.conclusion": "Duplicate customer keys cluster in a single source extract batch and need source-team confirmation.", "conclusion.confidence": 0.68, "diagnostics": "duplicate distribution"},
    ])
    approvals = pd.DataFrame([
        {"item_id": "appr_col_brand", "item_type": "column_mapping", "pair_id": "brand_dim", "table_name": "prod.analytics.dim_brand", "column_name": "brand_sid", "proposed_value": "brand_key -> brand_sid", "confidence": 0.74, "approval_status": "PENDING", "reviewer_comments": ""},
        {"item_id": "appr_rca_sales", "item_type": "rca_learning", "pair_id": "sales_fact", "table_name": "prod.analytics.fact_sales", "column_name": "", "proposed_value": "Late-arriving ERP filter mismatch pattern", "confidence": 0.82, "approval_status": "PENDING", "reviewer_comments": ""},
    ])
    summary = table_summary_for_ui(tables, results, rca, approvals)
    history = pd.DataFrame([
        {"run_id": "demo_20260629T090000Z", "status": "COMPLETED_WITH_FAILURES", "started_at": "2026-06-29T09:00:00Z", "completed_at": "2026-06-29T09:18:00Z", "tables": 4, "approval_proposals": 6},
        {"run_id": "demo_20260628T090000Z", "status": "COMPLETED", "started_at": "2026-06-28T09:00:00Z", "completed_at": "2026-06-28T09:14:00Z", "tables": 4, "approval_proposals": 1},
    ])
    return {
        "is_demo": True,
        "run_id": "demo_latest",
        "manifest": {"run_id": "demo_latest", "status": "DEMO_DATA", "started_at": "2026-06-29T09:00:00Z", "completed_at": "2026-06-29T09:18:00Z"},
        "history": history,
        "tables": tables,
        "results": results,
        "rca": rca,
        "approvals": approvals,
        "summary": summary,
        "report": {"Tables": tables, "RCA": rca, "Approvals": approvals},
        "events": pd.DataFrame([
            {"timestamp": "2026-06-29T09:01:00Z", "event": "run_started", "level": "INFO"},
            {"timestamp": "2026-06-29T09:12:00Z", "event": "rca_completed", "level": "INFO"},
        ]),
    }


def load_monitoring_snapshot(root: str | Path = ".", run_id: str | None = None) -> dict[str, Any]:
    project_root = root_path(root)
    history = load_run_history(project_root)
    selected = run_id or (None if history.empty else str(history.iloc[0]["run_id"]))
    if not selected:
        return _demo_snapshot()
    config = load_app_config(project_root)
    manifest = _read_json(config.path(config.project.logs_dir) / selected / "manifest.json")
    report = load_report_workbook(project_root, selected)
    tables = report.get("Tables", pd.DataFrame())
    if tables.empty and manifest.get("tables"):
        tables = pd.DataFrame(manifest.get("tables") or [])
    results = consolidated_results(report)
    rca = report.get("RCA", pd.DataFrame())
    approvals = report.get("Approvals", pd.DataFrame())
    summary = table_summary_for_ui(tables, results, rca, approvals)
    events = pd.DataFrame(read_jsonl(config.path(config.project.logs_dir) / selected / "events.jsonl", limit=50))
    return {
        "is_demo": False,
        "run_id": selected,
        "manifest": manifest,
        "history": history,
        "tables": tables,
        "results": results,
        "rca": rca,
        "approvals": approvals,
        "summary": summary,
        "report": report,
        "events": events,
    }


def dashboard_metrics(snapshot: dict[str, Any]) -> dict[str, int | str]:
    summary = snapshot.get("summary", pd.DataFrame())
    results = snapshot.get("results", pd.DataFrame())
    approvals = snapshot.get("approvals", pd.DataFrame())
    statuses = results.get("status", pd.Series(dtype=object)).astype(str).str.upper() if not results.empty else pd.Series(dtype=object)
    return {
        "Total tables validated": len(summary),
        "Passed tables": int((summary.get("Overall status", pd.Series(dtype=object)) == "Passed").sum()) if not summary.empty else 0,
        "Failed tables": int((summary.get("Overall status", pd.Series(dtype=object)) == "Failed").sum()) if not summary.empty else 0,
        "Tables with warnings": int((summary.get("Overall status", pd.Series(dtype=object)).isin(["Warning", "Approval Required"])).sum()) if not summary.empty else 0,
        "Total checks executed": len(results),
        "Failed checks": int(statuses.isin(["FAIL", "ERROR"]).sum()),
        "Checks requiring approval": len(approvals),
        "Latest validation run status": snapshot.get("manifest", {}).get("status", "No run"),
    }


def issue_counts(snapshot: dict[str, Any]) -> pd.DataFrame:
    results = snapshot.get("results", pd.DataFrame())
    if results.empty or "status" not in results:
        return pd.DataFrame(columns=["Category", "Issues"])
    statuses = results["status"].astype(str).str.upper()
    failed = results[statuses.isin(["FAIL", "ERROR"])]
    if failed.empty:
        return pd.DataFrame(columns=["Category", "Issues"])
    counts = failed.apply(lambda row: ui_category(row.to_dict()), axis=1).value_counts()
    return pd.DataFrame({"Category": counts.index, "Issues": counts.values})


def rca_completion(snapshot: dict[str, Any]) -> pd.DataFrame:
    summary = snapshot.get("summary", pd.DataFrame())
    if summary.empty:
        return pd.DataFrame(columns=["RCA status", "Tables"])
    counts = summary["RCA status"].value_counts()
    return pd.DataFrame({"RCA status": counts.index, "Tables": counts.values})


def run_progress(root: str | Path, run_id: str) -> dict[str, Any]:
    config = load_app_config(root_path(root))
    run_dir = config.path(config.project.logs_dir) / run_id
    manifest = _read_json(run_dir / "manifest.json")
    checkpoint = _read_json(run_dir / "checkpoint.json")
    events = read_jsonl(run_dir / "events.jsonl", limit=25)
    return {"manifest": manifest, "checkpoint": checkpoint, "events": events, "run_dir": run_dir}


def load_context_records(root: str | Path = ".") -> pd.DataFrame:
    config = load_app_config(root_path(root))
    store = make_context_retriever(config)
    store.sync()
    rows = []
    for record in store.get_exact():
        row = dict(record)
        row["payload"] = _json_text(row.get("payload"))
        row["provenance"] = _json_text(row.get("provenance"))
        rows.append(row)
    return pd.DataFrame(rows)


def context_asset_summaries(root: str | Path = ".") -> dict[str, pd.DataFrame]:
    config = load_app_config(root_path(root))
    relationships = read_yaml(config.path(config.project.relationships), default={})
    measures = read_yaml(config.path(config.project.measures), default={"measures": []})
    dimensions = []
    for name, payload in (relationships.get("dimensions") or {}).items():
        dimensions.append({
            "dimension": name,
            "enabled": bool((payload or {}).get("enabled", True)),
            "table": (payload or {}).get("table"),
            "description": (payload or {}).get("description"),
        })
    custom = []
    for item in relationships.get("custom_relationships") or []:
        custom.append({
            "id": item.get("id"), "enabled": item.get("enabled", True),
            "child_table": item.get("child_table"), "child_columns": _list_text(item.get("child_columns")),
            "parent_table": item.get("parent_table"), "parent_columns": _list_text(item.get("parent_columns")),
        })
    measure_rows = []
    for item in measures.get("measures") or []:
        measure_rows.append({
            "measure_id": item.get("measure_id"), "pair_id": item.get("pair_id"),
            "business_name": item.get("business_name"), "definition_type": item.get("definition_type"),
            "aggregation": item.get("aggregation"), "enabled": item.get("enabled", True),
            "approval_status": item.get("approval_status"),
        })
    return {
        "dimensions": pd.DataFrame(dimensions),
        "custom_relationships": pd.DataFrame(custom),
        "measures": pd.DataFrame(measure_rows),
    }


def add_approved_context_record(root: str | Path, record: dict[str, Any]) -> int:
    config = load_app_config(root_path(root))
    payload = record.get("payload")
    if isinstance(payload, str):
        payload = json.loads(payload) if payload.strip() else {}
    item = {
        "context_type": record.get("context_type") or "manual_context",
        "subject_key": record.get("subject_key") or stable_id(utc_now(), payload),
        "payload": payload or {},
        "pair_id": record.get("pair_id") or None,
        "target_table": record.get("target_table") or None,
        "source_table": record.get("source_table") or None,
        "column_name": record.get("column_name") or None,
        "business_entity": record.get("business_entity") or None,
        "table_type": record.get("table_type") or None,
        "confidence": float(record.get("confidence") or 1.0),
        "reviewed_by": record.get("reviewed_by") or "ui_reviewer",
        "reviewer_comments": record.get("reviewer_comments") or "Added from Streamlit UI",
    }
    return make_context_retriever(config).upsert_approved([item])


def list_approval_workbooks(root: str | Path = ".", state: str = "pending") -> list[Path]:
    config = load_app_config(root_path(root))
    folder = config.path(config.project.approvals.root) / state
    return sorted(folder.glob("*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)


def normalize_approval_frame(frame: pd.DataFrame) -> pd.DataFrame:
    output = _ensure_columns(frame, APPROVAL_COLUMNS)
    output = output.where(pd.notna(output), "")
    output["approval_status"] = output["approval_status"].astype(str).str.strip().str.upper().replace("", "PENDING")
    invalid = sorted(set(output["approval_status"]) - APPROVAL_STATUSES)
    if invalid:
        raise ValueError(f"Invalid approval statuses: {invalid}")
    return output


def save_approval_frame(path: Path, frame: pd.DataFrame) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    output = normalize_approval_frame(frame)
    output.to_excel(path, index=False, sheet_name="Approvals")
    workbook = load_workbook(path)
    sheet = workbook["Approvals"]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    validation = DataValidation(type="list", formula1='"PENDING,APPROVE,REJECT,NEEDS_CHANGES,OVERRIDE"', allow_blank=False)
    sheet.add_data_validation(validation)
    validation.add(f"L2:L{max(sheet.max_row, 2)}")
    workbook.save(path)
    return path


def publish_approval_decisions(root: str | Path, source_path: Path, frame: pd.DataFrame) -> dict[str, Any]:
    project_root = root_path(root)
    config = load_app_config(project_root)
    output = normalize_approval_frame(frame)
    if (output["approval_status"] == "PENDING").any():
        raise ValueError("Resolve every pending row before publishing approvals.")
    reviewed_dir = config.path(config.project.approvals.root) / "reviewed"
    reviewed_dir.mkdir(parents=True, exist_ok=True)
    reviewed_path = reviewed_dir / source_path.name
    if reviewed_path.exists():
        reviewed_path = reviewed_dir / f"{source_path.stem}_{stable_id(utc_now(), length=8)}{source_path.suffix}"
    save_approval_frame(reviewed_path, output)
    result = process_approval_workbook(config, reviewed_path)
    if source_path.exists() and source_path.parent.name == "pending":
        archive = config.path(config.project.approvals.root) / "archive"
        archive.mkdir(parents=True, exist_ok=True)
        archived = archive / f"{source_path.stem}_published_{stable_id(utc_now(), length=8)}{source_path.suffix}"
        shutil.move(str(source_path), str(archived))
        result["archived_pending_file"] = str(archived)
    return result


def pending_approval_frame(path: Path) -> pd.DataFrame:
    return read_approval_workbook(path)


def approved_context_from_frame(frame: pd.DataFrame) -> list[dict[str, Any]]:
    return approval_rows_to_context(normalize_approval_frame(frame))
